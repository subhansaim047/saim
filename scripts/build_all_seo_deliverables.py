import csv
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print("Generating Master SEO Suite...")

OUTPUT_DIR = os.path.join(os.getcwd(), "audit")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Helper function to auto-format worksheet
def format_sheet(ws, title, subtitle):
    # Insert Title Block
    ws.insert_rows(1, 3)
    ws['A1'] = title
    ws['A1'].font = Font(name="Calibri", size=15, bold=True, color="1F2937")
    ws['A2'] = subtitle
    ws['A2'].font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    # Format header row (now at row 4)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=4, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Format data rows
    for row in range(5, ws.max_row + 1):
        is_alt = (row % 2 == 0)
        row_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid") if is_alt else PatternFill(fill_type=None)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            if not cell.fill.start_color.rgb:
                cell.fill = row_fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10, color="1F2937")
            if col in [1, 2, 7, 8, 9, 10]:
                cell.alignment = Alignment(vertical="center")
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                
    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[3:]: # start from header
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

print("Base helper functions ready.")
